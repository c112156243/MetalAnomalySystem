"""真實 PySide6 事件迴圈 + mock SMTP；不連線、不使用真實認證。"""

import json
import os
import secrets
import smtplib
import ssl
import tempfile
import threading
import time
import unittest
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path
from openpyxl import load_workbook
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QTimer, qInstallMessageHandler
from PySide6.QtWidgets import QApplication

import main
import daq_acquisition as daq
import email_notifier as mail


class EmailTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config_path = Path(self.temp.name) / "config.json"
        self.config = mail.apply_email_defaults({
            "gateway1": "192.0.2.1", "gateway2": "192.0.2.2",
            "save_path": str(Path(self.temp.name) / "data"),
            "account": " first@example.com ", "device_number": "TEST裝置",
            "custom": {"preserve": True},
        })
        self.config["email"].update(enabled=True, username="sender@example.com",
                                      sender="sender@example.com")
        self.config_path.write_text(json.dumps(self.config), encoding="utf-8")
        self.patch(main, "CONFIG_FILE", str(self.config_path))
        self.patch(main, "show_message_box")
        self.daq_read = self.patch(daq, "acquire_voltage", return_value=self.sample_capture())
        self.env = patch.dict(os.environ, {"METAL_SMTP_PASSWORD": secrets.token_urlsafe(24)})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.client = Mock()
        self.client.send_message.return_value = {}
        self.ssl_factory = self.patch(mail.smtplib, "SMTP_SSL", return_value=self.client)
        self.tls_factory = self.patch(mail.smtplib, "SMTP", return_value=self.client)
        self.window = main.MainWindow()
        self.window.ui.show()
        self.addCleanup(self.close_window)

    def patch(self, target, name, *args, **kwargs):
        patcher = patch.object(target, name, *args, **kwargs)
        value = patcher.start()
        self.addCleanup(patcher.stop)
        return value

    def close_window(self):
        if self.window._capture_worker is not None:
            self.window._capture_worker.shutdown()
            self.window._capture_worker.wait()
            self.window.poll_capture_results()
        self.window.notifier.shutdown()
        self.window.notifier.wait()
        self.window.ui.close()
        self.app.processEvents()

    def until(self, predicate, timeout=4):
        end = time.monotonic() + timeout
        while time.monotonic() < end:
            self.app.processEvents()
            if predicate():
                return
            time.sleep(0.005)
        self.fail("等待 Qt／背景寄送狀態逾時")

    def sample_capture(self, when=None):
        return daq.CaptureResult(
            tuple(index / 1000 for index in range(1000)),
            "Dev1/ai0", 1000.0, "DIFF", -10.0, 10.0,
            "USB-6211", False, when or datetime(2026, 9, 24, 12, 0),
        )

    def capture(self, score=0.12, prediction="NG", when=None):
        with patch.object(self.window, "predict", return_value=(score, prediction)):
            self.window.complete_capture(self.sample_capture(when))

    def finished(self, text="已交給郵件伺服器"):
        self.until(lambda: text in self.window.email_status_label.text())

    def test_ok_has_no_job(self):
        self.capture(0.8, "OK")
        self.assertIsNone(self.window.notifier._thread)
        self.ssl_factory.assert_not_called()
        self.assertIn("本次 OK", self.window.email_status_label.text())

    def test_ng_one_message_utf8_and_no_attachment(self):
        when = datetime(2026, 9, 17, 12, 34, 56)
        self.capture(0.126, when=when)
        self.finished()
        self.assertEqual(self.client.send_message.call_count, 1)
        message = self.client.send_message.call_args.args[0]
        decoded = BytesParser(policy=policy.default).parsebytes(message.as_bytes())
        self.assertEqual(str(decoded["Subject"]), "【NG 異常通知】裝置 TEST裝置")
        self.assertEqual(str(decoded["To"]), "first@example.com")
        self.assertIn("擷取時間：2026-09-17 12:34", decoded.get_content())
        self.assertIn("預測分數：0.13", decoded.get_content())
        self.assertIn("預測結果：NG", decoded.get_content())
        self.assertEqual(list(decoded.iter_attachments()), [])
        self.assertEqual(self.ssl_factory.call_args.kwargs["timeout"], 20)
        self.assertTrue(self.ssl_factory.call_args.kwargs["context"].check_hostname)
        self.assertEqual(self.ssl_factory.call_args.kwargs["context"].verify_mode, ssl.CERT_REQUIRED)

    def test_consecutive_ng_each_sends_once_same_worker(self):
        self.capture()
        worker = self.window.notifier._thread
        self.capture(0.23)
        self.finished()
        self.assertEqual(self.client.send_message.call_count, 2)
        self.assertIs(worker, self.window.notifier._thread)
        self.assertEqual(self.window.tableHistory.rowCount(), 2)

    def test_save_excel_and_add_history_do_not_send(self):
        self.capture()
        self.finished()
        self.window.save_data()
        self.window.add_history()
        self.assertEqual(self.client.send_message.call_count, 1)
        workbook = load_workbook(self.window.last_excel_path)
        try:
            rows = list(workbook["擷取資料"].values)
            self.assertEqual(len(rows), 1001)
            self.assertEqual(list(rows[0]), ["Timestamp", "Voltage", "Current", "Score", "Prediction"])
            self.assertEqual(rows[1][-2:], (0.12, "NG"))
            self.assertIsInstance(rows[1][0], datetime)
            self.assertEqual([row[1] for row in rows[1:]], self.window.current_voltage)
            self.assertTrue(all(row[2] is None for row in rows[1:]))
            self.assertIn("尚未串接 AI", workbook["擷取設定"]["B14"].value)
        finally:
            workbook.close()

    def test_settings_dialog_save_reload_next_recipient(self):
        self.capture()
        self.finished()

        def save_dialog(dialog):
            dialog.account_edit.setText(" second@example.com ")
            dialog.save_settings()
            return dialog.result()

        with patch.object(main.SettingsDialog, "exec", save_dialog):
            self.window.open_settings()
        self.capture()
        self.finished()
        recipients = [call.kwargs["to_addrs"] for call in self.client.send_message.call_args_list]
        self.assertEqual(recipients, [["first@example.com"], ["second@example.com"]])
        saved = json.loads(self.config_path.read_text(encoding="utf-8"))
        self.assertEqual(saved["custom"], {"preserve": True})
        self.assertEqual(saved["email"], self.config["email"])

    def test_inflight_snapshot_and_old_result_do_not_overwrite_latest(self):
        started, release = threading.Event(), threading.Event()
        seen_threads = []

        def delayed_send(message, **kwargs):
            seen_threads.append(threading.get_ident())
            started.set()
            release.wait(3)
            return {}

        self.client.send_message.side_effect = delayed_send
        self.addCleanup(release.set)
        self.capture(0.11, when=datetime(2026, 9, 17, 1, 2))
        self.until(started.is_set)
        self.window.config["account"] = "second@example.com"
        self.window.config["device_number"] = "SECOND"
        self.window.config["email"]["sender"] = "new@example.com"
        self.window.config["email"]["smtp_host"] = "smtp.example.com"
        self.capture(0.22, when=datetime(2026, 9, 17, 3, 4))
        self.window.config["account"] = "third@example.com"
        self.window.config["email"]["sender"] = "third-sender@example.com"
        self.window.config["email"]["smtp_host"] = "third.example.com"
        self.capture(0.9, "OK")
        latest = self.window.email_status_label.text()
        self.window.status_label.setText("進行中：正在儲存資料...")
        release.set()
        self.until(lambda: "已交給" in self.window.tableHistory.item(1, 2).toolTip())
        calls = self.client.send_message.call_args_list
        first, second = [call.args[0] for call in calls]
        self.assertEqual(str(first["To"]), "first@example.com")
        self.assertEqual(str(first["From"]), "sender@example.com")
        self.assertIn("TEST裝置", str(first["Subject"]))
        self.assertIn("01:02", first.get_content())
        self.assertIn("0.11", first.get_content())
        self.assertEqual(str(second["To"]), "second@example.com")
        self.assertEqual(str(second["From"]), "new@example.com")
        self.assertIn("03:04", second.get_content())
        self.assertIn("0.22", second.get_content())
        self.assertEqual(self.ssl_factory.call_args_list[0].args[0], "smtp.gmail.com")
        self.assertEqual(self.ssl_factory.call_args_list[1].args[0], "smtp.example.com")
        self.assertEqual(self.window.email_status_label.text(), latest)
        self.assertEqual(self.window.status_label.text(), "進行中：正在儲存資料...")
        self.assertTrue(all(value != threading.get_ident() for value in seen_threads))

    def test_legacy_and_partial_config_preserve_values(self):
        for block in (None, {"smtp_host": "smtp.example.com", "custom": 42}):
            data = {key: value for key, value in self.config.items() if key != "email"}
            if block is not None:
                data["email"] = block
            self.config_path.write_text(json.dumps(data), encoding="utf-8")
            loaded = main.load_config()
            self.assertTrue(set(mail.EMAIL_DEFAULTS).issubset(loaded["email"]))
            self.assertEqual(loaded["account"], data["account"])
            self.assertEqual(loaded["custom"], data["custom"])
            if block is not None:
                self.assertEqual(loaded["email"]["smtp_host"], "smtp.example.com")
                self.assertEqual(loaded["email"]["custom"], 42)
            self.window.config = loaded
            self.capture()
            self.assertIn("通知未啟用", self.window.email_status_label.text())
            legacy_window = main.MainWindow()
            try:
                self.assertIn("通知未啟用", legacy_window.email_status_label.text())
            finally:
                legacy_window.ui.close()
        self.ssl_factory.assert_not_called()

    def test_missing_config_creates_disabled_defaults(self):
        self.config_path.unlink()
        fresh_window = main.MainWindow()
        try:
            saved = json.loads(self.config_path.read_text(encoding="utf-8"))
            self.assertEqual(saved["email"], mail.EMAIL_DEFAULTS)
            self.assertIn("通知未啟用", fresh_window.email_status_label.text())
            self.assertIsNone(fresh_window.notifier._thread)
        finally:
            fresh_window.ui.close()
        template = json.loads((Path(main.BASE_DIR) / "config.example.json").read_text(encoding="utf-8"))
        self.assertEqual(template["email"], mail.EMAIL_DEFAULTS)
        self.ssl_factory.assert_not_called()

    def test_invalid_settings_and_missing_password_preserve_capture(self):
        variants = [
            ("account", ""), ("account", "wrong"), ("account", "a@b\r\nBcc: x@y.com"),
            ("email", None), ("email", {"enabled": True}),
        ]
        for key, value in variants:
            with self.subTest(key=key, value=value):
                self.window.config = {**self.config, key: value}
                self.capture()
                self.assertIn("設定錯誤", self.window.email_status_label.text())
                self.assertEqual(self.window.current_prediction, "NG")
        for key, value in [("smtp_host", ""), ("username", ""), ("security", "plain"),
                           ("smtp_port", 0), ("timeout_seconds", 0),
                           ("timeout_seconds", float("inf")), ("enabled", "true")]:
            self.window.config = {**self.config, "email": {**self.config["email"], key: value}}
            self.capture()
            self.assertIn("設定錯誤", self.window.email_status_label.text())
        self.window.config = self.config
        with patch.dict(os.environ, {}, clear=True):
            self.capture()
            self.assertIn("密碼環境變數", self.window.email_status_label.text())
        self.ssl_factory.assert_not_called()

    def test_smtp_failures_do_not_lose_history_or_excel_and_never_retry(self):
        errors = [smtplib.SMTPAuthenticationError(535, b"redacted"), TimeoutError(),
                  smtplib.SMTPRecipientsRefused({}), ConnectionRefusedError(),
                  ssl.SSLError(), smtplib.SMTPDataError(554, b"redacted"), RuntimeError()]
        for index, error in enumerate(errors):
            with self.subTest(error=type(error).__name__):
                self.client.login.side_effect = error if index == 0 else None
                self.client.send_message.side_effect = error if index != 0 else None
                self.capture()
                self.finished("寄送失敗")
                self.assertEqual(self.window.tableHistory.rowCount(), index + 1)
                self.assertEqual(self.window.current_prediction, "NG")
                self.window.save_data()
                self.assertTrue(Path(self.window.last_excel_path).exists())
        self.assertEqual(self.client.login.call_count, len(errors))

    def test_refused_recipient_return_and_cleanup_after_acceptance(self):
        self.client.send_message.return_value = {"first@example.com": (550, b"rejected")}
        self.capture()
        self.finished("拒絕收件地址")
        self.client.send_message.return_value = {}
        self.client.close.side_effect = OSError()
        self.capture()
        self.finished()
        self.assertEqual(self.client.send_message.call_count, 2)

    def test_smtp_disconnect_and_rejection_have_specific_reasons(self):
        for error, reason in [
            (smtplib.SMTPServerDisconnected(), "中斷連線，寄送結果可能不明"),
            (smtplib.SMTPDataError(554, b"rejected"), "拒絕寄送或通訊失敗"),
        ]:
            with self.subTest(error=type(error).__name__):
                self.client.send_message.side_effect = error
                self.capture()
                self.finished(reason)
                self.assertNotIn("無法連線", self.window.email_status_label.text())
        self.assertEqual(self.client.send_message.call_count, 2)

    def test_completion_during_capture_keeps_capture_status(self):
        self.capture()
        self.finished()
        self.window.capture_data()
        label = self.window.email_status_label.text()
        self.window.update_email_status(1, "寄送失敗：測試")
        self.assertEqual(self.window.email_status_label.text(), label)
        self.assertIn("正在擷取", self.window.status_label.text())

    def test_starttls_before_login_and_no_plaintext_fallback(self):
        self.window.config["email"]["security"] = "starttls"
        self.window.config["email"]["smtp_port"] = 587
        self.capture()
        self.finished()
        names = [call[0] for call in self.client.method_calls]
        self.assertEqual(names[:5], ["ehlo", "starttls", "ehlo", "login", "send_message"])
        self.assertTrue(self.client.starttls.call_args.kwargs["context"].check_hostname)
        self.client.reset_mock()
        self.client.starttls.side_effect = ssl.SSLError()
        self.capture()
        self.finished("TLS 加密")
        self.client.login.assert_not_called()
        self.client.send_message.assert_not_called()
        self.ssl_factory.assert_not_called()

    def test_close_active_send_cancels_pending_and_keeps_event_loop_alive(self):
        started, release = threading.Event(), threading.Event()
        warnings = []
        previous = qInstallMessageHandler(lambda kind, context, message: warnings.append(message))
        self.addCleanup(qInstallMessageHandler, previous)

        def delayed(*args, **kwargs):
            started.set()
            release.wait(3)
            return {}

        self.client.send_message.side_effect = delayed
        self.addCleanup(release.set)
        self.capture()
        self.until(started.is_set)
        self.capture(0.22)
        self.window.btnExit.click()
        self.assertTrue(self.window.ui.isVisible())
        self.assertFalse(self.window.btnCapture.isEnabled())
        ticks = []
        QTimer.singleShot(0, lambda: ticks.append(True))
        self.until(lambda: bool(ticks))
        self.until(lambda: "已取消" in self.window.tableHistory.item(1, 2).toolTip())
        release.set()
        self.until(lambda: not self.window.ui.isVisible())
        self.assertFalse(self.window.notifier.is_running())
        self.assertEqual(self.client.send_message.call_count, 1)
        self.assertFalse(any("Destroyed while" in message or "still running" in message for message in warnings))

    def test_capture_button_worker_and_close_during_capture(self):
        with patch.object(self.window, "predict", return_value=(0.1, "NG")):
            self.window.btnCapture.click()
            self.assertFalse(self.window.btnCapture.isEnabled())
            self.until(lambda: self.window.tableHistory.rowCount() == 1)
            self.finished()
        self.window.btnCapture.click()
        self.window.ui.close()
        self.until(lambda: not self.window.ui.isVisible())
        self.assertFalse(self.window._capture_timer.isActive())
        self.assertEqual(self.window.tableHistory.rowCount(), 1)

    def test_bounded_queue_reports_overflow(self):
        started, release = threading.Event(), threading.Event()

        def delayed(*args, **kwargs):
            started.set()
            release.wait(3)
            return {}

        self.client.send_message.side_effect = delayed
        self.addCleanup(release.set)
        self.capture()
        self.until(started.is_set)
        for _ in range(101):
            self.capture()
        self.assertIn("佇列已滿", self.window.email_status_label.text())
        self.window.notifier.shutdown()
        release.set()

    def test_daq_failure_keeps_previous_batch_without_prediction_or_email(self):
        self.capture()
        self.finished()
        previous = self.window.current_capture
        self.daq_read.side_effect = daq.AcquisitionError("NI-DAQmx 錯誤 -50103：設備被占用")
        with patch.object(self.window, "predict") as predict:
            self.window.btnCapture.click()
            self.until(lambda: self.window._capture_worker is None)
            predict.assert_not_called()
        self.assertIs(self.window.current_capture, previous)
        self.assertIn("保留上一批", self.window.status_label.text())
        self.assertEqual(self.window.tableHistory.rowCount(), 1)
        self.assertEqual(self.client.send_message.call_count, 1)
        self.assertTrue(self.window.btnCapture.isEnabled())
        self.assertTrue(self.window.btnSave.isEnabled())
        self.assertTrue(self.window.btnSettings.isEnabled())
        self.daq_read.side_effect = None
        with patch.object(self.window, "predict", return_value=(0.8, "OK")):
            self.window.btnCapture.click()
            self.until(lambda: self.window.tableHistory.rowCount() == 2)
        self.assertEqual(self.client.send_message.call_count, 1)

    def test_close_slow_capture_cancels_without_late_result_or_mail(self):
        started = threading.Event()
        ticks = []
        seen_threads = []

        def delayed(settings, cancel):
            seen_threads.append(threading.get_ident())
            started.set()
            cancel.wait(3)
            # 模擬取消與完成競爭：即使回傳成功，也不能在關窗後判定／寄信。
            return self.sample_capture()

        self.daq_read.side_effect = delayed
        self.window.btnCapture.click()
        self.until(started.is_set)
        self.window.capture_data()
        self.assertEqual(self.daq_read.call_count, 1)
        self.assertFalse(self.window.btnSave.isEnabled())
        self.assertFalse(self.window.btnSettings.isEnabled())
        QTimer.singleShot(0, lambda: ticks.append(True))
        self.until(lambda: bool(ticks))
        self.window.ui.close()
        self.until(lambda: not self.window.ui.isVisible())
        self.assertIsNone(self.window._capture_worker)
        self.assertEqual(self.window.tableHistory.rowCount(), 0)
        self.ssl_factory.assert_not_called()
        self.assertNotEqual(seen_threads[0], threading.get_ident())

    def test_invalid_daq_config_no_read_and_legacy_defaults_preserved(self):
        self.assertEqual(self.window.config["daq"], daq.DAQ_DEFAULTS)
        self.window.config["daq"]["sample_rate"] = 0
        self.window.btnCapture.click()
        self.daq_read.assert_not_called()
        self.assertIn("擷取失敗", self.window.status_label.text())
        self.ssl_factory.assert_not_called()
        partial = daq.apply_daq_defaults({"daq": {"channel": "Dev2/ai1", "custom": 42}})
        self.assertEqual(partial["daq"]["channel"], "Dev2/ai1")
        self.assertEqual(partial["daq"]["custom"], 42)
        self.assertEqual(partial["daq"]["sample_rate"], 1000)

    def test_excel_repeat_save_uses_unique_paths_and_same_capture_metadata(self):
        self.capture(0.8, "OK")
        with patch.object(main, "datetime") as clock:
            clock.now.return_value = datetime(2026, 9, 24, 15, 0)
            self.window.save_data()
            first = Path(self.window.last_excel_path)
            initial_bytes = first.read_bytes()
            self.window.config["daq"]["channel"] = "Dev2/ai1"
            self.window.save_data()
            second = Path(self.window.last_excel_path)
        self.assertNotEqual(first, second)
        self.assertEqual(first.read_bytes(), initial_bytes)
        self.assertTrue(second.name.endswith("_002.xlsx"))
        workbook = load_workbook(second)
        try:
            self.assertEqual(workbook["擷取設定"]["B2"].value, "Dev1/ai0")
            self.assertEqual(workbook["擷取設定"]["B8"].value, 1000)
        finally:
            workbook.close()
        self.ssl_factory.assert_not_called()

    def test_excel_save_failure_keeps_data_and_restores_button(self):
        self.capture(0.8, "OK")
        previous = self.window.current_capture
        for target in ("create_today_folder", "save_capture_excel"):
            with self.subTest(target=target):
                with patch.object(main, target, side_effect=PermissionError("test")):
                    self.window.save_data()
                self.assertIn("儲存失敗", self.window.status_label.text())
                self.assertIs(self.window.current_capture, previous)
                self.assertTrue(self.window.btnSave.isEnabled())
        self.assertIsNone(self.window.last_excel_path)

    def test_save_without_capture_does_not_create_file(self):
        self.window.save_data()
        self.assertIsNone(self.window.last_excel_path)
        self.assertEqual(list(Path(self.temp.name).rglob("*.xlsx")), [])

    def test_logs_do_not_include_exception_payload_or_secret(self):
        payload = os.environ["METAL_SMTP_PASSWORD"]
        self.client.login.side_effect = smtplib.SMTPAuthenticationError(535, payload.encode())
        with self.assertLogs(mail.logger, level="WARNING") as logs:
            self.capture()
            self.finished("寄送失敗")
        self.assertNotIn(payload, "\n".join(logs.output))
        self.assertNotIn(payload, self.window.email_status_label.text())
        self.assertIn("SMTPAuthenticationError", "\n".join(logs.output))


if __name__ == "__main__":
    unittest.main()
