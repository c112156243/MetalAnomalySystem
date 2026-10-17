"""Excel 輸出：保留實測電壓與程式模擬電流，分別註明來源。"""

from io import BytesIO
from pathlib import Path


def save_capture_excel(folder, capture, score, prediction, saved_at, segments=None, report_png=None):
    if len(capture.voltage) != len(capture.current):
        raise ValueError("電壓與模擬電流的樣本數不一致，無法儲存。")
    if len(capture.voltage) < 1024:
        raise ValueError("資料未足1024筆，無法儲存。")
    if segments is not None and len(segments) != len(capture.voltage) // 1024:
        raise ValueError("分段判定數量與擷取資料不一致。")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "擷取資料"
    sheet.append(["Timestamp", "Voltage", "Current", "Score", "Prediction", "CurrentSource"])
    for index, (voltage, current) in enumerate(zip(capture.voltage, capture.current)):
        row_score, row_prediction = score, prediction
        if segments is not None:
            segment = segments[index // 1024] if index // 1024 < len(segments) else None
            row_score = segment.score if segment else None
            row_prediction = segment.prediction if segment else "未判定（不足1024筆）"
        sheet.append([saved_at, voltage, current, row_score, row_prediction, "程式模擬（非實測）"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 23, "B": 20, "C": 18, "D": 14, "E": 18, "F": 28}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"
        row[1].number_format = "0.000000"
        row[3].number_format = "0.00"
        row[2].number_format = "0.000000"

    metadata = workbook.create_sheet("擷取設定")
    for row in [
        ["項目", "內容"],
        ["通道", capture.channel],
        ["裝置型號", capture.product_type],
        ["電壓來源", "NI-DAQmx 模擬裝置" if capture.is_simulated else "實體裝置"],
        ["端點模式", capture.terminal_config],
        ["設定電壓下限 (V)", capture.min_voltage],
        ["設定電壓上限 (V)", capture.max_voltage],
        ["實際取樣率 (samples/s)", capture.sample_rate],
        ["樣本數", len(capture.voltage)],
        ["擷取完成時間", capture.completed_at],
        ["儲存時間", saved_at],
        ["Voltage 單位", "V"],
        ["Current", "程式模擬（非實測），不由電壓換算"],
        ["Score / Prediction", "隨機測試判定，尚未串接 AI 模型"],
        ["Timestamp", "儲存時間，非逐筆取樣時間"],
        ["Current 單位", "A"],
        ["電流模擬波形", "正弦波"],
        ["電流模擬峰值 (A)", capture.current_simulation.amplitude_a],
        ["電流模擬頻率 (Hz)", capture.current_simulation.frequency_hz],
        ["電流模擬偏移 (A)", capture.current_simulation.offset_a],
        ["時間對應", "同列為相同樣本索引；時間為索引 / 實際取樣率，電流由程式產生"],
    ]:
        metadata.append(row)
        # 設定中的文字一律以字串儲存，避免被 Excel 視為公式。
        for cell in metadata[metadata.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif hasattr(cell.value, "strftime"):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
    metadata.column_dimensions["A"].width = 32
    metadata.column_dimensions["B"].width = 52
    metadata.freeze_panes = "A2"
    if segments is not None:
        metadata.append(["辨識窗長", 1024])
        metadata.append(["未判定尾段筆數", len(capture.voltage) % 1024])
        metadata.append(["整次結果", prediction])
        metadata.append(["整次分數（最低段分數）", score])
        metadata.append(["OK 比例", sum(item.prediction == "OK" for item in segments) / len(segments)])
        results = workbook.create_sheet("分段判定")
        results.append(["段號", "起始樣本索引", "結束樣本索引", "Score", "Prediction", "來源"])
        for index, item in enumerate(segments, 1):
            results.append([index, item.start, item.stop - 1, item.score, item.prediction, "隨機模擬，非 AI"])
        results.freeze_panes = "A2"
        results.auto_filter.ref = results.dimensions
        for column in "ABCDEF":
            results.column_dimensions[column].width = 22
    for table in (sheet, metadata):
        for cell in table[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")

    buffer = BytesIO()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    root = Path(folder)
    root.mkdir(parents=True, exist_ok=True)
    prefix = saved_at.strftime("%Y%m%d_%H%M%S")
    index = 1
    while True:
        path = root / f"{prefix}_{index:03d}.xlsx"
        png_path = path.with_suffix(".png")
        try:
            stream = path.open("xb")
        except FileExistsError:
            index += 1
            continue
        png_stream = None
        if report_png is not None:
            try:
                png_stream = png_path.open("xb")
            except FileExistsError:
                stream.close()
                path.unlink()
                index += 1
                continue
            except Exception:
                stream.close()
                path.unlink(missing_ok=True)
                raise
        try:
            with stream:
                stream.write(buffer.getvalue())
            if png_stream is not None:
                with png_stream:
                    png_stream.write(report_png)
        except Exception:
            if png_stream is not None:
                png_stream.close()
                png_path.unlink(missing_ok=True)
            path.unlink(missing_ok=True)
            raise
        return str(path)
